# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

NAVY  = RGBColor(0x1F, 0x3B, 0x5B)   # X 自变量
TEAL  = RGBColor(0x1B, 0x7F, 0x79)   # M 中介机制
GREEN = RGBColor(0x2E, 0x7D, 0x32)   # Y 良性结果/纠偏
RED   = RGBColor(0xB0, 0x3A, 0x2E)   # Y 风险结果
ORANGE= RGBColor(0xC0, 0x5A, 0x16)   # 真实风险
PURPLE= RGBColor(0x6A, 0x3D, 0x9A)   # 调节W / 表达成本
AMBER  = RGBColor(0xC8, 0x7F, 0x0A)  # 识别/冲击
LAMBER = RGBColor(0xF1, 0xD9, 0xA8)
GREY  = RGBColor(0x6B, 0x76, 0x82)   # 控制Z
BLUE  = RGBColor(0x2E, 0x5E, 0x8C)   # 技术路线步骤
LGREY = RGBColor(0xEC, 0xEF, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK  = RGBColor(0x20, 0x28, 0x33)
LINEC = RGBColor(0x5A, 0x6B, 0x7B)

SW, SH = 13.333, 7.5
FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width  = Inches(SW)
prs.slide_height = Inches(SH)


def new_slide(title):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                            Inches(SW), Inches(0.62))
    sp.fill.solid(); sp.fill.fore_color.rgb = NAVY
    sp.line.fill.background(); sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title
    r.font.size = Pt(18); r.font.bold = True
    r.font.color.rgb = WHITE; r.font.name = FONT
    return s


def box(s, x, y, w, h, text, fill, fg=WHITE, size=11, bold=True,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = line if line else fill
    sp.line.width = Pt(1); sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    first = True
    for ln in text.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = fg; r.font.name = FONT
    return sp


def label(s, x, y, w, h, text, size=10, color=DARK, italic=False,
          bold=False, align=PP_ALIGN.CENTER):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    first = True
    for ln in text.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.color.rgb = color
        r.font.italic = italic; r.font.bold = bold; r.font.name = FONT
    return tb


def conn(s, p1, p2, color=LINEC, width=1.5, dashed=False, arrow=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                               Inches(p1[0]), Inches(p1[1]),
                               Inches(p2[0]), Inches(p2[1]))
    c.line.color.rgb = color; c.line.width = Pt(width)
    c.shadow.inherit = False
    ln = c.line._get_or_add_ln()
    if dashed:
        ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'dash'}))
    if arrow:
        ln.append(ln.makeelement(qn('a:tailEnd'),
                  {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    return c


def bc(s):
    return (s.left + s.width / 2) / 914400, (s.top + s.height) / 914400


def tc(s):
    return (s.left + s.width / 2) / 914400, s.top / 914400


def lc(s):
    return s.left / 914400, (s.top + s.height / 2) / 914400


def rc(s):
    return (s.left + s.width) / 914400, (s.top + s.height / 2) / 914400


def htag(s, x, y, txt, color, w=0.55):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                            Inches(w), Inches(0.30))
    sp.fill.solid(); sp.fill.fore_color.rgb = WHITE
    sp.line.color.rgb = color; sp.line.width = Pt(1.25)
    sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(1); tf.margin_right = Pt(1)
    tf.margin_top = Pt(0); tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = txt
    r.font.size = Pt(10); r.font.bold = True
    r.font.color.rgb = color; r.font.name = FONT
    return sp


def legend(s, items, y=7.05):
    x = 0.3
    for txt, col in items:
        d = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                               Inches(0.22), Inches(0.16))
        d.fill.solid(); d.fill.fore_color.rgb = col
        d.line.fill.background(); d.shadow.inherit = False
        label(s, x + 0.26, y - 0.06, 1.9, 0.3, txt, 9, DARK,
              align=PP_ALIGN.LEFT)
        x += 0.30 + 0.10 + min(1.9, 0.16 * len(txt) + 0.6)


def footer(s, text, y=6.62):
    box(s, 0.3, y, SW - 0.6, 0.42, text, NAVY, WHITE, 10.5, True,
        MSO_SHAPE.RECTANGLE)


def rolelab(s, y, t1, color):
    label(s, 0.16, y, 1.55, 0.55, t1, 11, color, bold=True)


# ============================================================ Slide 1
s = new_slide("总框架·三章串联：风险产生 → 风险显化 → 风险纠偏（闭环）")
box(s, 0.3, 0.70, SW - 0.6, 0.44,
    "总问题：LLM 触达升级是否在抬高获客转化的同时推高借款人脆弱性与市场风险，"
    "监管与智慧投诉治理又能否纠偏？", BLUE, WHITE, 11.5, True, MSO_SHAPE.RECTANGLE)

CX = 2.97; CW = 7.4

def chip(x, y, t, col):
    sp = box(s, x, y, 1.5, 0.46, t, WHITE, col, 10.5, True)
    sp.line.color.rgb = col; sp.line.width = Pt(1.5)
    return sp

X1 = box(s, CX, 1.26, CW, 0.50,
         "X｜LLM 触达升级：电销外呼 IVR → LLM 大模型接入", NAVY, size=12)
M1 = box(s, CX, 1.84, CW, 0.46,
         "M｜转化：发起用信申请（意图）→ 成功用信（转化）", TEAL, size=11)
Y1 = box(s, CX, 2.38, CW, 0.58,
         "Y₁｜机构内借款人金融脆弱性：复借·逾期(DPD30+)·额度使用率·借后多头",
         RED, size=11)
label(s, CX, 2.99, CW, 0.22, "▼ 机构内风险在市场显化为投诉", 9.5, GREY,
      italic=True)
RR = box(s, 3.05, 3.22, 3.55, 0.54, "真实风险暴露 ↑", ORANGE, size=11)
EE = box(s, 6.72, 3.22, 3.65, 0.54,
         "客户表达成本 ↓（投诉文本长度/规范性↑）", PURPLE, size=10)
Y2 = box(s, CX, 3.86, CW, 0.62,
         "Y₂｜互联网消金高风险投诉 ↑\n催收/逾期/征信/高息/诱导/信息骚扰；占全部投诉比",
         RED, size=10.5)
label(s, CX, 4.51, CW, 0.22, "▼ 可观测风险成为治理对象", 9.5, GREY,
      italic=True)
X3 = box(s, CX, 4.74, CW, 0.50,
         "X₃｜2025 智慧315 / 监管治理冲击", AMBER, size=12)
M3 = box(s, CX, 5.30, CW, 0.46,
         "M｜治理过程：回复率·解决率·响应时长·监管处罚", TEAL, size=11)
Y3 = box(s, CX, 5.84, CW, 0.56,
         "Y₃｜治理结果：投诉总数↓·重复投诉↓·高风险占比↓·未解决↓",
         GREEN, size=11)
chip(0.40, 2.38, "Q1·产生", RED)
chip(0.40, 3.22, "Q2·显化", ORANGE)
chip(0.40, 4.74, "Q3·纠偏", GREEN)
conn(s, bc(X1), tc(M1)); conn(s, bc(M1), tc(Y1))
conn(s, bc(Y1), tc(RR)); conn(s, bc(Y1), tc(EE))
conn(s, bc(RR), (tc(Y2)[0]-0.3, tc(Y2)[1]))
conn(s, bc(EE), (tc(Y2)[0]+0.3, tc(Y2)[1]))
conn(s, bc(Y2), tc(X3)); conn(s, bc(X3), tc(M3)); conn(s, bc(M3), tc(Y3))
conn(s, (rc(Y3)[0]+0.05, rc(Y3)[1]), (rc(Y2)[0]+0.05, rc(Y2)[1]),
     GREEN, 1.3, dashed=True)
label(s, 10.5, 4.4, 2.6, 0.3, "纠偏闭环 ↩", 9.5, GREEN, italic=True,
      align=PP_ALIGN.LEFT)
footer(s, "结论与政策闭环：金融AI消费者保护 · 智能外呼合规 · 投诉治理 — "
          "同一风险构念，三套数据三测互证", 6.56)
s.notes_slide.notes_text_frame.text = (
    "1. 第一章在机构内造出脆弱性（X→转化→Y1）；"
    "2. 第二章在市场显化为投诉，并分解真实风险 vs 表达成本；"
    "3. 第三章治理冲击经治理过程纠偏，回拉投诉、闭环；"
    "同一风险构念被机构数据/公开投诉/治理准实验三测互证。")

# ============================================================ Slide 2
s = new_slide("第一章 框架图（风险产生）  X →（中介M·转化）→ Y，调节W·控制Z")
rolelab(s, 1.30, "自变量\nX", NAVY)
rolelab(s, 2.70, "中介\nM", TEAL)
rolelab(s, 1.95, "因变量\nY", RED)
rolelab(s, 4.55, "调节\nW", PURPLE)
LX = 1.85
X  = box(s, LX, 1.15, 4.3, 0.85,
         "X｜LLM 外呼接入·触达升级\n剂量梯度：IVR → LLM-v1 → LLM-v2", NAVY, size=11)
M  = box(s, LX, 2.55, 4.3, 0.85,
         "M｜转化（中介，非落脚）\n① 是否发起用信申请(意图)  ② 是否成功用信(转化)",
         TEAL, size=10.5)
Y  = box(s, 7.05, 1.55, 5.9, 1.05,
         "Y｜客户金融脆弱性\n复借次数·逾期率(DPD30+)·额度使用率\n借后下载其他借贷App(applist多头)·脆弱性指数",
         RED, size=11)
DM = box(s, 5.55, 2.62, 1.35, 0.72, "外延 vs\n内涵(H3)", AMBER, size=9.5,
         shape=MSO_SHAPE.DIAMOND)
W  = box(s, LX, 4.40, 4.3, 0.85,
         "调节 W（H4·异质）\nIVR→LLM 智能化梯度（剂量响应）· 客户风险类型/征信分",
         PURPLE, size=10)
conn(s, bc(X), tc(M))
conn(s, rc(M), lc(DM)); conn(s, rc(DM), (lc(Y)[0], lc(Y)[1]+0.1))
label(s, 1.95, 2.02, 4.1, 0.3, "更多人接触消金（外延边际）", 9, GREY,
      italic=True)
conn(s, tc(W), (rc(M)[0]-0.6, bc(M)[1]+0.02), PURPLE, 1.4, dashed=True)
htag(s, 5.9, 1.62, "H1", NAVY)
htag(s, 6.45, 2.78, "H2", RED)
htag(s, 5.78, 3.38, "H3", AMBER)
htag(s, 5.45, 3.95, "H4", PURPLE)
footer(s, "控制 Z：征信分·历史多头·负债收入比·年龄·名单来源·活动类型·时段·"
          "地域　|　识别：以 LLM 接入日为外生时点的 Stacked RD-in-time")
legend(s, [("自变量X", NAVY), ("中介M", TEAL), ("因变量Y", RED),
           ("调节W", PURPLE), ("分解H3", AMBER)])

# ============================================================ Slide 3
s = new_slide("第一章 技术路线图（风险产生）  问题→数据→识别→分解→检验→结论")
P  = box(s, 3.5, 0.80, 6.3, 0.54,
         "研究问题：LLM 接入是否经转化把更多更脆弱的人卷入并推高其脆弱性？",
         BLUE, size=11.5)
D  = box(s, 3.5, 1.52, 6.3, 0.50,
         "数据：全量AI拨打 + 每通模型版本 + 用信表 + applist多头", BLUE, size=11)
G  = box(s, 3.5, 2.20, 6.3, 0.50,
         "反推 LLM 接入日 g  +  采用曲线诊断（陡跳 / 灰度爬坡）", BLUE, size=11)
R1 = box(s, 1.5, 2.98, 4.6, 0.56, "陡跳 → Stacked RD-in-time（主模型）", TEAL, size=10.5)
R2 = box(s, 7.2, 2.98, 4.6, 0.56, "爬坡 → Fuzzy RD / 剂量交错DiD", AMBER, size=10.5)
EST= box(s, 3.5, 3.74, 6.3, 0.52,
         "估计：中介 M＝转化(意图/用信)  与  Y＝客户金融脆弱性", GREEN, size=11)
DEC= box(s, 3.5, 4.44, 6.3, 0.56,
         "分解：外延(经转化卷入更脆弱人群) vs 内涵(同类恶化)\n因果中介 / DFL重加权",
         AMBER, size=10)
TT = box(s, 3.5, 5.22, 6.3, 0.50,
         "识别检验：密度连续 · 安慰剂日 · donut-RD · 机制三角", GREY, size=10.5)
conn(s, bc(P), tc(D)); conn(s, bc(D), tc(G))
conn(s, bc(G), tc(R1)); conn(s, bc(G), tc(R2))
conn(s, bc(R1), (tc(EST)[0]-0.3, tc(EST)[1]))
conn(s, bc(R2), (tc(EST)[0]+0.3, tc(EST)[1]))
conn(s, bc(EST), tc(DEC)); conn(s, bc(DEC), tc(TT))
footer(s, "结论：LLM 接入经转化推高借款人脆弱性，且“被更多卷入”（外延）为主驱动")

# ============================================================ Slide 4
s = new_slide("第二章 框架图（风险显化）  X → 双中介 → Y，核心识别 H6")
rolelab(s, 1.05, "自变量\nX", NAVY)
rolelab(s, 2.95, "双中介\n（命门）", PURPLE)
rolelab(s, 1.95, "因变量\nY", RED)
rolelab(s, 4.95, "调节\nW", PURPLE)
LX = 1.85
X2 = box(s, LX, 1.05, 4.2, 0.62,
         "X｜LLM 上线\n（机构级错时上线）", NAVY, size=11)
RR = box(s, LX, 2.45, 4.2, 0.66,
         "真实风险暴露 ↑\n诱导借款·高息·暴力催收等真实损害", ORANGE, size=10)
EE = box(s, LX, 3.25, 4.2, 0.66,
         "客户表达成本 ↓\nLLM 降低投诉撰写门槛：文本长度/规范性/要素完备度↑",
         PURPLE, size=9.5)
Y2 = box(s, 7.1, 1.95, 5.85, 1.30,
         "Y｜互联网消金高风险投诉 ↑\n催收·逾期·征信·高息/隐形收费·诱导借款·"
         "个人信息骚扰（数量）\n互联网消金投诉数 / 全部投诉数（占比）",
         RED, size=10.5)
W2 = box(s, LX, 4.55, 4.2, 0.66,
         "调节 W｜是否持牌机构 · 上线前机构高风险投诉占比", PURPLE, size=10)
conn(s, bc(X2), (tc(RR)[0], tc(RR)[1]-0.06))
conn(s, rc(RR), (lc(Y2)[0], lc(Y2)[1]-0.18))
conn(s, rc(EE), (lc(Y2)[0], lc(Y2)[1]+0.18))
conn(s, (rc(RR)[0]-0.1, bc(RR)[1]), (rc(EE)[0]-0.1, tc(EE)[1]),
     PURPLE, 1.2, dashed=True, arrow=False)
conn(s, tc(W2), (lc(Y2)[0], lc(Y2)[1]+0.42), PURPLE, 1.3, dashed=True)
htag(s, 6.18, 1.05, "H5", RED)
htag(s, 6.18, 2.78, "H6", AMBER, w=1.55)
label(s, 5.95, 3.18, 2.0, 0.3, "核心识别：分解二者", 8.5, AMBER,
      italic=True, align=PP_ALIGN.LEFT)
htag(s, 6.18, 4.55, "H7", PURPLE)
footer(s, "识别：机构级错时 DiD / 事件研究　|　核心：以投诉文本特征为表达成本"
          "代理，分离“净真实风险”与“表达便利虚增”")
legend(s, [("自变量X", NAVY), ("真实风险", ORANGE), ("表达成本/调节", PURPLE),
           ("因变量Y", RED), ("分解H6", AMBER)])

# ============================================================ Slide 5
s = new_slide("第二章 技术路线图（风险显化）  问题→数据→识别→真实vs表达分解")
P  = box(s, 3.0, 0.82, 7.3, 0.62,
         "研究问题：LLM 上线是否放大互金消金风险？投诉上升中真实风险 vs "
         "表达成本各几何？", BLUE, size=11)
D  = box(s, 3.0, 1.62, 7.3, 0.58,
         "数据：黑猫投诉/12378（备选裁判文书）+ 机构LLM上线时点 + 投诉文本特征",
         BLUE, size=10.5)
M  = box(s, 3.0, 2.40, 7.3, 0.58,
         "主识别：机构级错时 DiD / 事件研究（稳健估计量 + 前趋势检验）", TEAL, size=10.5)
KEY= box(s, 3.0, 3.18, 7.3, 0.74,
         "核心分解（H6）：文本特征构造表达成本代理 E\n"
         "C = π·Post + φ·E + ψ·(Post×E)；控制 E 后 π ＝ 净真实风险显化",
         AMBER, size=10)
PL = box(s, 3.0, 4.12, 7.3, 0.58,
         "安慰剂：非消金投诉作对照（表达成本同步、真实风险不应同步）+ 可模板化子类",
         GREY, size=10)
R  = box(s, 3.0, 4.90, 7.3, 0.56,
         "结论：净真实风险显化幅度 与 表达便利虚增 的分离估计", ORANGE, size=11)
conn(s, bc(P), tc(D)); conn(s, bc(D), tc(M))
conn(s, bc(M), tc(KEY)); conn(s, bc(KEY), tc(PL)); conn(s, bc(PL), tc(R))
footer(s, "要点：解决“LLM 也让人更会写投诉”的长期混淆——投诉量↑不等于真实风险↑")

# ============================================================ Slide 6
s = new_slide("第三章 框架图（风险纠偏）  准实验 · X →（治理过程M）→ Y")
rolelab(s, 1.00, "识别\n准实验", AMBER)
rolelab(s, 2.45, "自变量\nX", NAVY)
rolelab(s, 3.70, "中介\nM", TEAL)
rolelab(s, 4.95, "因变量\nY", GREEN)
LX = 1.85
DZ = box(s, LX, 0.95, 9.6, 0.56,
         "准实验：处理组 持牌 vs 非持牌（或上线前高风险占比 高 vs 低）× 前后  "
         "DiD / 三重差分", AMBER, size=10.5)
X3 = box(s, LX, 2.35, 9.6, 0.56,
         "X｜2025 智慧315 / 监管治理冲击（Post）", NAVY, size=11)
M3 = box(s, LX, 3.60, 9.6, 0.56,
         "M｜治理过程：投诉回复率↑ · 解决率↑ · 平均响应时长↓ · 是否伴随监管处罚",
         TEAL, size=10.5)
Y3 = box(s, LX, 4.85, 9.6, 0.62,
         "Y｜治理结果：投诉总数↓ · 重复投诉↓ · 高风险投诉占比↓ · 未解决投诉↓",
         GREEN, size=11)
conn(s, tc(X3), bc(DZ), AMBER, 1.3, dashed=True)
conn(s, bc(X3), tc(M3)); conn(s, bc(M3), tc(Y3))
htag(s, 6.4, 3.05, "H9", TEAL)
htag(s, 6.4, 4.30, "H8", GREEN)
footer(s, "稳健：PSM-DID · 替代处理组定义 · 事件研究前趋势　|　局限：2025 冲击"
          "数据窗口偏薄，以高频月度弥补")
legend(s, [("识别/冲击", AMBER), ("自变量X", NAVY), ("中介M", TEAL),
           ("因变量Y", GREEN)])

# ============================================================ Slide 7
s = new_slide("第三章 技术路线图（风险纠偏）  问题→数据→DiD→中介→稳健→政策")
P  = box(s, 3.2, 0.86, 6.9, 0.62,
         "研究问题：监管与智慧投诉治理能否、经由何种过程缓解 LLM 带来的消金风险？",
         BLUE, size=11)
D  = box(s, 3.2, 1.66, 6.9, 0.56,
         "数据：第三方投诉 + 监管处罚 + 2025智慧315治理时点（机构×月面板）",
         BLUE, size=10.5)
M  = box(s, 3.2, 2.42, 6.9, 0.56,
         "主识别：处理组 vs 对照组 × 前后  DiD / 三重差分", TEAL, size=11)
MED= box(s, 3.2, 3.18, 6.9, 0.62,
         "治理过程中介检验：回复率/解决率/响应时长/处罚 是否为纠偏路径\n"
         "（对照“投诉自然回落”安慰剂）", AMBER, size=10)
ROB= box(s, 3.2, 4.10, 6.9, 0.56,
         "稳健：PSM-DID · 替代处理组 · 事件研究前趋势 ·（窗口偏薄→高频月度补）",
         GREY, size=10)
R  = box(s, 3.2, 4.86, 6.9, 0.58,
         "结论与政策闭环：金融AI消保 · 智能外呼合规 · 投诉治理", ORANGE, size=11)
conn(s, bc(P), tc(D)); conn(s, bc(D), tc(M)); conn(s, bc(M), tc(MED))
conn(s, bc(MED), tc(ROB)); conn(s, bc(ROB), tc(R))
footer(s, "三章闭环：产生（机构内）→ 显化（市场投诉，真实/表达两分）→ 纠偏"
          "（治理冲击经治理过程回拉）")

out = "/home/user/claude/技术路线图_LLM消费信贷.pptx"
prs.save(out)
print("SAVED", out, "slides=", len(prs.slides._sldIdLst))
